#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""解析抖音官方三表，按区间对齐作品，生成采集清单与数据底稿。

输入（周期目录内，缺失时自动到系统下载目录寻找最新同名文件）：
  - 作品列表导出.xlsx       （作品粒度，全量，本地按区间筛）
  - 数据表现_短视频全量指标.xlsx（日粒度，平台按选定区间导出）
  - 数据表现_粉丝全量指标.xlsx （日粒度，可选）

输出（写入周期目录）：
  - _采集清单.json  结构化数据 + 每条作品的本地文件对齐状态（供后续脚本/Agent 使用）
  - _数据底稿.md    数据表格 + 衍生指标（Agent 写归因报告时的数据依据，不做主观结论）

兼容性：三表的百分比/时长/整数在不同表里格式不同（作品表是小数 0.118、
日表是 "11.81%"、时长有 "4.13s"/"63.46秒"/纯数字、缺失为 "-"），统一在此归一。

用法：
  python parse_data.py --start 2026-10-09 --end 2026-10-15
  python parse_data.py --start 2026-10-01 --end 2026-10-16 --offline   # 只扫本地文件
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    daterange, load_ledger, save_ledger, materials_root, get_downloads_dir,
    parse_date, period_dir_name, build_file_base, probe_media,
)

try:
    import openpyxl
except ImportError:
    sys.exit("缺少 openpyxl，请在对应 Python 环境执行：pip install openpyxl")


# ── 单元格归一 ────────────────────────────────────────────────────────────────

_BLANK = {"-", "—", "--", "", "None", "nan"}


def _num(v):
    """从任意单元格提取 float；空值返回 None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("，", "")
    if s in _BLANK:
        return None
    # 取首个数字片段（兼容 "4.13s"、"63.46秒"、"≈123"）
    import re
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def to_pct(v):
    """归一为比率小数。"11.81%" -> 0.1181；"0.118" -> 0.118；缺失 -> None。"""
    if v is None:
        return None
    s = str(v).strip()
    if s in _BLANK:
        return None
    n = _num(v)
    if n is None:
        return None
    if "%" in s:
        return n / 100.0
    return n


def to_int(v):
    n = _num(v)
    return int(round(n)) if n is not None else None


def to_seconds(v):
    """归一为秒。"4.13s"、"63.46秒"、"4.77" -> float 秒。"""
    return _num(v)


def parse_dt(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def pct_str(x, digits=2):
    return "-" if x is None else f"{x * 100:.{digits}f}%"


def num_str(x, digits=2):
    return "-" if x is None else f"{x:.{digits}f}"


# ── 工作簿定位与读取 ─────────────────────────────────────────────────────────

def latest_match(keyword: str, period_dir: Path, mode: str = "period") -> Path | None:
    """定位数据文件。

    - mode="full"（作品列表，全量导出）：周期目录与下载目录合并，取修改时间最新的。
      全量表越新越全，本地会按区间筛，因此即使来自下载目录也没关系。
    - mode="period"（日表/粉丝表，按区间导出）：优先周期目录内已归档的文件，
      因为它才对应本周期；周期目录没有才去下载目录兜底。
    """
    def _scan(d: Path) -> list[Path]:
        if not d.is_dir():
            return []
        return [f for f in d.glob("*.xlsx")
                if keyword in f.name and not f.name.startswith("~$")]

    local = _scan(period_dir)
    downloaded = _scan(get_downloads_dir())
    if mode == "full":
        pool = local + downloaded
        return max(pool, key=lambda f: f.stat().st_mtime) if pool else None
    if local:
        return max(local, key=lambda f: f.stat().st_mtime)
    if downloaded:
        return max(downloaded, key=lambda f: f.stat().st_mtime)
    return None


def read_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() in _BLANK for c in r):
            continue
        out.append({header[i]: r[i] for i in range(len(header))})
    return out


# ── 三表标准化 ───────────────────────────────────────────────────────────────

def parse_works(path: Path, start: date, end: date) -> list[dict]:
    works = []
    for r in read_rows(path):
        dt = parse_dt(r.get("发布时间"))
        if dt is None or not (start <= dt.date() <= end):
            continue
        plays = to_int(r.get("播放量"))
        likes = to_int(r.get("点赞量"))
        shares = to_int(r.get("分享量"))
        comments = to_int(r.get("评论量"))
        collects = to_int(r.get("收藏量"))
        engage_components = [x for x in (likes, shares, comments, collects) if x is not None]
        engagement = (sum(engage_components) / plays) if plays else None
        works.append({
            "name": str(r.get("作品名称") or "").strip(),
            "publish_time": dt.strftime("%Y-%m-%d %H:%M:%S"),
            "_dt": dt,
            "genre": str(r.get("体裁") or "").strip(),
            "status": str(r.get("审核状态") or "").strip(),
            "plays": plays,
            "completion_rate": to_pct(r.get("完播率")),
            "complete_5s": to_pct(r.get("5s完播率")),
            "cover_ctr": to_pct(r.get("封面点击率")),
            "bounce_2s": to_pct(r.get("2s跳出率")),
            "avg_play_sec": to_seconds(r.get("平均播放时长")),
            "likes": likes, "shares": shares, "comments": comments,
            "collects": collects,
            "homepage_visits": to_int(r.get("主页访问量")),
            "new_fans": to_int(r.get("粉丝增量")),
            "engagement": engagement,
        })
    works.sort(key=lambda w: w["_dt"])
    return works


def parse_daily(path: Path, start: date, end: date) -> dict[str, dict]:
    daily = {}
    for r in read_rows(path):
        d = parse_date(r.get("日期")) if r.get("日期") is not None else None
        if d is None or not (start <= d <= end):
            continue
        plays = to_int(r.get("总播放量"))
        likes = to_int(r.get("总点赞量"))
        shares = to_int(r.get("总分享量"))
        comments = to_int(r.get("总评论量"))
        comps = [x for x in (likes, shares, comments) if x is not None]
        daily[d.isoformat()] = {
            "date": d.isoformat(), "posts": to_int(r.get("投稿量")),
            "plays": plays, "likes": likes, "shares": shares, "comments": comments,
            "complete_5s": to_pct(r.get("5秒完播率")),
            "bounce_2s": to_pct(r.get("2秒跳出率")),
            "cover_ctr": to_pct(r.get("封面点击率")),
            "avg_play_sec": to_seconds(r.get("平均播放时长")),
            "engagement": (sum(comps) / plays) if plays else None,
        }
    return daily


def parse_fans(path: Path, start: date, end: date) -> dict[str, dict]:
    fans = {}
    for r in read_rows(path):
        d = parse_date(r.get("日期")) if r.get("日期") is not None else None
        if d is None or not (start <= d <= end):
            continue
        fans[d.isoformat()] = {
            "date": d.isoformat(), "total_fans": to_int(r.get("总粉丝量")),
            "net_fans": to_int(r.get("粉丝净增")), "gained": to_int(r.get("吸粉量")),
            "lost": to_int(r.get("脱粉量")), "returning": to_int(r.get("回访粉丝量")),
        }
    return fans


# ── 离线扫描（无三表，仅本地素材）────────────────────────────────────────────

def scan_offline(period_dir: Path, start: date, end: date) -> list[dict]:
    works = []
    for mp4 in sorted(period_dir.glob("*.mp4")):
        # 文件名形如 20261011-标题 或 20261007-01-标题
        stem = mp4.stem
        d = None
        if len(stem) >= 8 and stem[:8].isdigit():
            d = datetime.strptime(stem[:8], "%Y%m%d").date()
        if d is None or not (start <= d <= end):
            continue
        works.append({
            "name": stem.split("-", 1)[-1], "publish_time": f"{d.isoformat()} 00:00:00",
            "_dt": datetime(d.year, d.month, d.day), "genre": "", "status": "本地文件",
            "plays": None, "completion_rate": None, "complete_5s": None, "cover_ctr": None,
            "bounce_2s": None, "avg_play_sec": None, "likes": None, "shares": None,
            "comments": None, "collects": None, "homepage_visits": None, "new_fans": None,
            "engagement": None,
        })
    return works


# ── 本地文件对齐 ─────────────────────────────────────────────────────────────

def attach_local_files(period_dir: Path, works: list[dict]) -> None:
    """为每条作品补标准文件主名、本地 mp4/mp3/txt/帧状态与实际视频时长。

    匹配策略：优先精确匹配标准名（新流程由 normalize 生成）；找不到则按
    “日期前缀 + 当天序号”回退匹配，兼容历史手动命名（20261011-单点突破法.mp4）。
    mp3/txt 跟随实际匹配到的 mp4 主名（历史素材三者同名）。
    """
    from datetime import datetime as _dt
    # 收集目录内所有 mp4，按 日期(前8位) 分组并排序
    mp4_by_day: dict = defaultdict(list)
    for mp4 in sorted(period_dir.glob("*.mp4")):
        if len(mp4.stem) >= 8 and mp4.stem[:8].isdigit():
            try:
                d = _dt.strptime(mp4.stem[:8], "%Y%m%d").date()
            except ValueError:
                continue
            mp4_by_day[d].append(mp4)

    per_day = defaultdict(int)
    for w in works:
        per_day[w["_dt"].date()] += 1

    for w in works:
        day = w["_dt"].date()
        count = per_day[day]
        idx = 0 if count == 1 else sum(
            1 for x in works if x["_dt"].date() == day and x["_dt"] <= w["_dt"])
        base = build_file_base(w["_dt"], idx, w["name"])
        w["index_in_day"] = idx
        w["file_base"] = base

        # 1) 精确匹配标准名
        mp4 = period_dir / f"{base}.mp4"
        # 2) 回退：按日期+序号
        if not mp4.exists():
            group = mp4_by_day.get(day, [])
            if count == 1 and len(group) >= 1:
                mp4 = group[0]
            elif count > 1 and len(group) >= idx:
                # 优先文件名带两位序号，否则按排序位置对应
                numbered = [p for p in group if len(p.stem) > 10 and p.stem[9:11].isdigit()]
                if numbered:
                    wanted = f"{day.strftime('%Y%m%d')}-{idx:02d}"
                    hit = next((p for p in numbered if p.stem.startswith(wanted)), None)
                    if hit:
                        mp4 = hit
                if not mp4.exists() and idx - 1 < len(group):
                    mp4 = group[idx - 1]

        stem = mp4.stem if mp4.exists() else base
        mp3 = period_dir / f"{stem}.mp3"
        txt = period_dir / f"{stem}.txt"
        frames_dir = period_dir / "frames" / stem
        w["files"] = {
            "mp4": mp4.name if mp4.exists() else None,
            "mp3": mp3.name if mp3.exists() else None,
            "txt": txt.name if txt.exists() else None,
            "frames": frames_dir.exists() and bool(list(frames_dir.glob("*.jpg"))),
        }
        w["script_chars"] = len(txt.read_text(encoding="utf-8").strip()) if txt.exists() else 0
        if mp4.exists():
            info = probe_media(mp4)
            w["video_duration"] = info["duration"]
            w["resolution"] = f"{info['width']}x{info['height']}" if info["width"] else None
            w["has_audio"] = info["has_audio"]
        else:
            w["video_duration"] = None
            w["resolution"] = None
            w["has_audio"] = None
        if w.get("avg_play_sec") and w.get("video_duration"):
            w["play_duration_ratio"] = w["avg_play_sec"] / w["video_duration"]
        else:
            w["play_duration_ratio"] = None


# ── 数据底稿 ─────────────────────────────────────────────────────────────────

def build_dossier(period_dir: Path, start: date, end: date, works, daily, fans) -> str:
    L = []
    L.append(f"# 数据底稿 {period_dir_name(start, end)}（脚本自动生成，勿手改）\n")
    L.append("> 本文件仅做事实汇总与口径统一，供归因分析引用；不含主观结论。\n")

    # 日表
    L.append("## 一、每日数据表现\n")
    L.append("| 日期 | 投稿 | 播放 | 点赞 | 分享 | 评论 | 5s完播 | 2s跳出 | 封面点击 | 平均播放 | 互动率 |")
    L.append("|---|---|---|---|---|---|---|---|---|---|---|")
    for d in daterange(start, end):
        r = daily.get(d.isoformat())
        if r:
            L.append(f"| {d.strftime('%m-%d')} | {r['posts'] if r['posts'] is not None else '-'} | "
                     f"{r['plays']} | {r['likes']} | {r['shares']} | {r['comments']} | "
                     f"{pct_str(r['complete_5s'])} | {pct_str(r['bounce_2s'])} | "
                     f"{pct_str(r['cover_ctr'])} | {num_str(r['avg_play_sec'])}s | "
                     f"{pct_str(r['engagement'])} |")
        else:
            L.append(f"| {d.strftime('%m-%d')} | - | - | - | - | - | - | - | - | - | - |")

    # 粉丝
    if fans:
        L.append("\n## 二、每日粉丝数据\n")
        L.append("| 日期 | 总粉丝 | 净增 | 吸粉 | 脱粉 | 回访 |")
        L.append("|---|---|---|---|---|---|")
        for d in daterange(start, end):
            r = fans.get(d.isoformat())
            if r:
                L.append(f"| {d.strftime('%m-%d')} | {r['total_fans']} | {r['net_fans']} | "
                         f"{r['gained']} | {r['lost']} | {r['returning']} |")
    else:
        L.append("\n## 二、每日粉丝数据\n\n（本周期未提供粉丝表，或粉丝量未开放相关画像，跳过）\n")

    # 作品明细
    L.append("\n## 三、作品明细（区间内按发布时间排序）\n")
    L.append("| # | 发布时间 | 标题 | 播放 | 完播 | 5s完播 | 2s跳出 | 平均播放 | 时长 | 时长占比 | 赞 | 藏 | 评 | 转 | 互动率 | 涨粉 | 本地文件 |")
    L.append("|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|")
    for i, w in enumerate(works, 1):
        f = w["files"]
        ready = []
        ready.append("MP4" if f["mp4"] else "缺视频")
        ready.append("MP3" if f["mp3"] else "缺音频")
        ready.append("脚本" if f["txt"] else "缺脚本")
        ready.append("帧" if f["frames"] else "缺帧")
        title = w["name"].split("#")[0][:16]
        L.append(
            f"| {i} | {w['publish_time'][5:16]} | {title} | {w['plays']} | "
            f"{pct_str(w['completion_rate'])} | {pct_str(w['complete_5s'])} | "
            f"{pct_str(w['bounce_2s'])} | {num_str(w['avg_play_sec'])}s | "
            f"{num_str(w.get('video_duration'))}s | {pct_str(w.get('play_duration_ratio'))} | "
            f"{w['likes']} | {w['collects']} | {w['comments']} | {w['shares']} | "
            f"{pct_str(w['engagement'])} | {w['new_fans']} | {'/'.join(ready)} |")

    # 汇总
    if works:
        tp = sum(w["plays"] for w in works if w["plays"] is not None)
        tl = sum(w["likes"] for w in works if w["likes"] is not None)
        tc = sum(w["collects"] for w in works if w["collects"] is not None)
        ts = sum(w["shares"] for w in works if w["shares"] is not None)
        tcm = sum(w["comments"] for w in works if w["comments"] is not None)
        tf = sum(w["new_fans"] for w in works if w["new_fans"] is not None)
        L.append("\n## 四、区间汇总\n")
        L.append(f"- 作品数：{len(works)} 条；总播放：{tp}；总赞：{tl}；总藏：{tc}；"
                 f"总评：{tcm}；总转：{ts}；总涨粉：{tf}")
        best = max((w for w in works if w["plays"] is not None), key=lambda x: x["plays"], default=None)
        if best:
            L.append(f"- 播放最高：{best['file_base']}（{best['plays']} 播放）")
        L.append("\n> 指标基准与口径见 references/METRICS_BENCHMARK.md；冷启动期注意熟人流量校准。")
    return "\n".join(L) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser(description="解析抖音三表并对齐本地素材，生成采集清单与数据底稿")
    ap.add_argument("--start", required=True, help="起始日期，如 2026-10-09")
    ap.add_argument("--end", required=True, help="结束日期，如 2026-10-15")
    ap.add_argument("--period-dir", default=None, help="周期目录（默认 素材根/YYYYMMDD-YYYYMMDD）")
    ap.add_argument("--offline", action="store_true", help="无三表，仅扫描本地 mp4 构建清单")
    args = ap.parse_args()

    start, end = parse_date(args.start), parse_date(args.end)
    if end < start:
        sys.exit("结束日期早于起始日期")
    period_dir = Path(args.period_dir) if args.period_dir else \
        materials_root() / period_dir_name(start, end)
    period_dir.mkdir(parents=True, exist_ok=True)
    print(f"周期目录：{period_dir}")

    daily, fans = {}, {}
    if args.offline:
        works = scan_offline(period_dir, start, end)
        print("offline 模式：仅扫描本地 mp4，无官方数据。")
    else:
        works_path = latest_match("作品列表", period_dir, mode="full")
        daily_path = latest_match("短视频", period_dir, mode="period")
        fans_path = latest_match("粉丝", period_dir, mode="period")
        if not works_path:
            sys.exit("未找到『作品列表导出.xlsx』；请先在作品管理页导出，或用 --offline。")
        print(f"作品表：{works_path.name}")
        print(f"日数据表：{daily_path.name if daily_path else '未找到（可仅有作品表）'}")
        print(f"粉丝表：{fans_path.name if fans_path else '未找到（可选）'}")
        works = parse_works(works_path, start, end)
        if daily_path:
            daily = parse_daily(daily_path, start, end)
        if fans_path:
            fans = parse_fans(fans_path, start, end)

    attach_local_files(period_dir, works)

    # 断更日检测（仅在有日表时）
    no_post = []
    if daily:
        for d in daterange(start, end):
            r = daily.get(d.isoformat())
            if r is None or not r.get("posts"):
                no_post.append(d.strftime("%m-%d"))
    multi_days = sorted({w["_dt"].date().strftime("%m-%d") for w in works
                         if w["index_in_day"] > 0})

    ledger = {
        "period": {"start": start.isoformat(), "end": end.isoformat(),
                   "dir": str(period_dir)},
        "works": [{k: v for k, v in w.items() if k != "_dt"} for w in works],
        "daily": list(daily.values()),
        "fans": list(fans.values()),
        "no_post_days": no_post,
        "multi_post_days": multi_days,
    }
    save_ledger(period_dir, ledger)
    dossier = build_dossier(period_dir, start, end, works, daily, fans)
    (period_dir / "_数据底稿.md").write_text(dossier, encoding="utf-8")

    # 核对清单
    print("\n" + "=" * 60)
    print(f"区间内作品：{len(works)} 条；日历日：{(end - start).days + 1} 天")
    if no_post:
        print(f"⚠ 断更/无投稿日：{', '.join(no_post)}")
    if multi_days:
        print(f"★ 一天多条：{', '.join(multi_days)}（已按 01/02 编号）")
    miss_mp4 = [w["file_base"] for w in works if not w["files"]["mp4"]]
    if miss_mp4:
        print(f"待下载视频 {len(miss_mp4)} 条：")
        for b in miss_mp4:
            print("   -", b)
    else:
        print("√ 全部视频已就位")
    print(f"已生成：_采集清单.json、_数据底稿.md")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    sys.exit(main())
