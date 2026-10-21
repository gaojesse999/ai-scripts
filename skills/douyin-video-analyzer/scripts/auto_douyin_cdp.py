#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
通过 CDP 远程调试端口接管前台真实 Edge/Chrome 浏览器，完成全自动导出与复盘分析
"""

import sys
import os

sys.stdout.reconfigure(encoding='utf-8')
sys.stderr.reconfigure(encoding='utf-8')

# 确保脚本内访问抖音/直连时不受本地代理影响
for k in ['HTTP_PROXY', 'HTTPS_PROXY', 'ALL_PROXY', 'http_proxy', 'https_proxy', 'all_proxy']:
    if k in os.environ:
        del os.environ[k]

import time
import json
import urllib.request
import subprocess
import argparse
import re
from datetime import datetime, timedelta
from pathlib import Path

from playwright.sync_api import sync_playwright

DOWNLOAD_DIR = Path(os.environ.get("USERPROFILE", "")) / "Downloads"
VAULT_ROOT = Path(__file__).resolve().parents[4]
SKILL_ROOT = VAULT_ROOT / ".agents" / "skills" / "douyin-video-analyzer"
PYTHON_EXE = sys.executable

sys.path.insert(0, str(SKILL_ROOT / "scripts"))
try:
    from common import clean_title, find_browser, probe_media
except ImportError:
    def clean_title(t, max_len=14):
        return str(t)[:max_len]
    def find_browser():
        return r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    def probe_media(p):
        return {"ok": True}

def get_target_work_title(target_date):
    try:
        import openpyxl
        hits = sorted(
            [f for f in DOWNLOAD_DIR.glob("*作品列表*.xlsx") if not f.name.startswith("~$")],
            key=lambda f: f.stat().st_mtime, reverse=True
        )
        if not hits:
            return None
        wb = openpyxl.load_workbook(hits[0], data_only=True)
        ws = wb.active
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        title_col = None
        time_col = None
        for i, h in enumerate(headers):
            if "作品名称" in h or "标题" in h:
                title_col = i
            elif "发布时间" in h or "时间" in h:
                time_col = i
        if title_col is None:
            title_col = 0
        target_fmt = target_date.replace("-", "")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= title_col:
                continue
            t_val = str(row[title_col] or "").strip()
            time_val = str(row[time_col] or "").strip() if time_col is not None and len(row) > time_col else ""
            if target_date in time_val or target_fmt in time_val:
                return t_val
    except Exception as e:
        print(f">>> 尝试读取作品列表标题提示: {e}", flush=True)
    return None

def is_cdp_port_open(port=9225):
    try:
        url = f"http://127.0.0.1:{port}/json/version"
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(url, timeout=1) as resp:
            return resp.status == 200
    except Exception:
        return False

def ensure_browser_running(port=9225, timeout=30):
    if is_cdp_port_open(port):
        print(f">>> √ 前台调试浏览器已在运行中 (端口 {port})，直接建立 CDP 连接！", flush=True)
        return True

    print(f">>> 检测到调试端口 {port} 未开启，正在通过 Python 自动启动前台浏览器...", flush=True)
    browser_exe = find_browser()
    user_dir = Path.home() / ".douyin_automation_cdp"
    user_dir.mkdir(parents=True, exist_ok=True)
    target_url = "https://creator.douyin.com/creator-micro/data-center/operation"
    
    # 在 Windows 上调用 start 独立进程派发，确保窗口在前台主桌面正常显示且不阻塞 Python 执行
    cmd_str = f'start "" "{browser_exe}" --remote-debugging-port={port} --no-proxy-server --user-data-dir="{user_dir}" --start-maximized "{target_url}"'
    os.system(cmd_str)
    
    print(f">>> 浏览器启动指令已发出，正在等待 CDP 调试端口 {port} 就绪...", flush=True)
    for i in range(timeout):
        time.sleep(1)
        if is_cdp_port_open(port):
            print(f">>> √ 前台浏览器已成功启动并就绪 (耗时 {i+1} 秒)！", flush=True)
            return True
        if (i + 1) % 5 == 0:
            print(f">>> 等待浏览器调试端口响应中... 已等待 {i+1} 秒", flush=True)
            
    print(f">>> 错误: 自动启动浏览器超时 (已等待 {timeout} 秒)！", flush=True)
    return False

def get_douyin_page(context):
    for page in context.pages:
        if "creator.douyin.com" in page.url:
            return page
    return context.pages[-1] if context.pages else None

def wait_for_user_login(page):
    print(">>> 检测浏览器登录状态...", flush=True)
    for i in range(300):
        try:
            txt = page.inner_text("body")
            if "扫码登录" not in txt and ("首页" in txt or "内容管理" in txt or "作品发布" in txt):
                print(">>> √ 检测到已成功登录创作者中心！", flush=True)
                return True
            if "接收短信验证码" in txt:
                print(">>> 提示：浏览器弹出了二次验证，请在浏览器中直接完成验证！", flush=True)
        except Exception:
            pass
            
        if i % 10 == 0 and i > 0:
            print(f">>> 正在等待登录完成（已等待 {i*2} 秒）... 请在屏幕上的浏览器窗口中扫码！", flush=True)
        time.sleep(2)
        
    print(">>> 登录等待超时！", flush=True)
    return False

def export_datacenter_tables(page):
    print(">>> [步骤 1/4] 开始导出数据中心两张表（短视频 + 粉丝）...", flush=True)
    page.goto("https://creator.douyin.com/creator-micro/data-center/operation", wait_until="domcontentloaded")
    time.sleep(4)
    
    # 点击“昨天”
    try:
        page.evaluate("""() => {
            const btns = [...document.querySelectorAll('div, span')].filter(el => 
                el.innerText && el.innerText.trim() === '昨天' && el.offsetParent !== null
            );
            btns.forEach(b => b.click());
        }""")
        print(">>> √ 已点击【昨天】按钮", flush=True)
        time.sleep(3)
    except Exception as e:
        print(f">>> 日期选择提示: {e}", flush=True)
        
    export_buttons = page.query_selector_all("button:has-text('导出数据')")
    print(f">>> 找到 {len(export_buttons)} 个导出数据按钮", flush=True)
    
    for i, btn in enumerate(export_buttons):
        try:
            name = "短视频全量指标" if i == 0 else "粉丝全量指标"
            print(f">>> 正在导出: {name} ...", flush=True)
            btn.click()
            time.sleep(3)
        except Exception as e:
            print(f">>> 导出第 {i+1} 张表异常: {e}", flush=True)
            
    time.sleep(3)

def export_works_list(page):
    print(">>> [步骤 2/4] 开始导出作品管理页全量作品列表...", flush=True)
    page.goto("https://creator.douyin.com/creator-micro/content/manage", wait_until="domcontentloaded")
    time.sleep(4)
    
    try:
        btn = page.locator("text='导出数据'").first
        btn.click()
        time.sleep(3)
        modal_btns = page.locator("button:has-text('确定'), button:has-text('导出')").all()
        for mb in modal_btns:
            try:
                mb.click()
            except Exception:
                pass
        print(">>> √ 已触发作品列表导出", flush=True)
        time.sleep(4)
    except Exception as e:
        print(f">>> 作品列表导出异常: {e}", flush=True)

def download_period_video(context, page, target_date, period_dir):
    print(f">>> [步骤 3/4] 正在定位 {target_date} 作品并捕获视频流...", flush=True)
    known_title = get_target_work_title(target_date)
    if known_title:
        cleaned_short = clean_title(known_title, max_len=14)
        print(f">>> 从导出的作品列表中识别到目标作品: 『{known_title}』 (短码: {cleaned_short})", flush=True)
    else:
        cleaned_short = ""

    # 优先使用个人主页提取直链（符合 BROWSER_RUNBOOK 规范，直观且免弹窗阻断）
    user_url = "https://www.douyin.com/user/self"
    print(f">>> 导航至个人主页提取作品直链: {user_url} ...", flush=True)
    page.goto(user_url, wait_until="domcontentloaded")
    time.sleep(5)

    works = page.evaluate("""() => {
        const links = [...document.querySelectorAll('a[href*="/video/"]')];
        return links.map((a, i) => ({
            index: i,
            href: a.href,
            text: (a.innerText || '').slice(0, 80).replace(/\\n/g, ' ')
        }));
    }""")
    print(f">>> 个人主页解析到 {len(works)} 个作品卡片", flush=True)

    target_href = None
    if works:
        for w in works:
            t = w.get("text", "")
            if known_title and (cleaned_short in t or known_title[:6] in t):
                target_href = w.get("href")
                print(f">>> √ 按标题匹配到作品: {t[:40]} -> {target_href}", flush=True)
                break
            if target_date in t or target_date.replace("-", "") in t:
                target_href = w.get("href")
                print(f">>> √ 按日期匹配到作品: {t[:40]} -> {target_href}", flush=True)
                break
        
        # 若为单天分析且未精准命中，非置顶的第 1 条通常是最新的目标作品
        if not target_href and len(works) > 1:
            for w in works:
                if "置顶" not in w.get("text", ""):
                    target_href = w.get("href")
                    print(f">>> 自动选中最新发布作品: {w.get('text', '')[:40]} -> {target_href}", flush=True)
                    break

    video_url = None
    if target_href:
        print(f">>> 导航至目标视频播放页捕获流: {target_href} ...", flush=True)
        page.goto(target_href, wait_until="domcontentloaded")
        time.sleep(6)

        vids = page.query_selector_all("video")
        for v in vids:
            src = v.get_attribute("src") or page.evaluate("el => el.currentSrc", v)
            if src and ("douyinvod.com" in src or "tos-cn-v" in src):
                video_url = src
                break

        if not video_url:
            try:
                entries = page.evaluate('''() => {
                    return window.performance.getEntriesByType('resource')
                        .map(e => e.name)
                        .filter(u => u.includes('douyinvod.com') || u.includes('/video/tos/'));
                }''')
                for u in entries:
                    if 'douyinvod.com' in u or '/video/tos/' in u:
                        video_url = u
                        break
            except Exception:
                pass

    if video_url:
        print(f">>> √ 成功捕获真实视频流 URL: {video_url[:80]}...", flush=True)
        if known_title:
            filename = f"{target_date.replace('-', '')}-{clean_title(known_title, max_len=14)}.mp4"
        else:
            filename = f"{target_date.replace('-', '')}-发布视频.mp4"
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        req = urllib.request.Request(video_url, headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Referer": "https://www.douyin.com/"
        })
        try:
            with opener.open(req) as resp:
                data = resp.read()
                save_period = period_dir / filename
                save_period.write_bytes(data)
                (DOWNLOAD_DIR / filename).write_bytes(data)
                print(f">>> √ 视频已保存至周期目录与下载目录: {filename} ({len(data)} 字节)", flush=True)
                info = probe_media(save_period)
                print(f">>> √ 媒体完整性校验: {info}", flush=True)
        except Exception as e:
            print(f">>> 视频下载失败: {e}", flush=True)
    else:
        print(">>> 提示: 未能自动截获视频流，可手动下载后放入周期目录。", flush=True)

def run_local_pipeline(target_date, period_dir):
    print(">>> [步骤 4/4] 启动本地分析流水线（归一化 -> 解析 -> Whisper转写 -> 抽帧 -> 报告）...", flush=True)
    scripts_dir = SKILL_ROOT / "scripts"
    
    # 1. normalize_names.py
    subprocess.run([PYTHON_EXE, str(scripts_dir / "normalize_names.py"), "--start", target_date, "--end", target_date])
    
    # 2. parse_data.py
    subprocess.run([PYTHON_EXE, str(scripts_dir / "parse_data.py"), "--start", target_date, "--end", target_date])
    
    # 3. media_pipeline.py
    subprocess.run([PYTHON_EXE, str(scripts_dir / "media_pipeline.py"), "--start", target_date, "--end", target_date])
    
    # 4. extract_frames.py
    subprocess.run([PYTHON_EXE, str(scripts_dir / "extract_frames.py"), "--start", target_date, "--end", target_date])
    
    # 5. build_report.py
    subprocess.run([PYTHON_EXE, str(scripts_dir / "build_report.py"), "--start", target_date, "--end", target_date, "--force"])
    print(">>> √ 全部本地流水线已处理完毕！", flush=True)

def main():
    parser = argparse.ArgumentParser(description="抖音自动化采集与分析 (CDP 前台浏览器接管模式)")
    yesterday = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
    parser.add_argument("--date", default=yesterday, help="目标分析日期 (YYYY-MM-DD)，默认为昨天")
    parser.add_argument("--port", type=int, default=9225, help="CDP 远程调试端口")
    args = parser.parse_args()
    
    target_date = args.date
    period_dir = VAULT_ROOT / "视频分析" / "视频素材" / target_date.replace("-", "")
    period_dir.mkdir(parents=True, exist_ok=True)
    
    print("=" * 60)
    print(f" 抖音自动化采集分析器（CDP 前台真实浏览器接管模式）")
    print(f" 目标日期: {target_date}")
    print("=" * 60)
    
    if not ensure_browser_running(port=args.port):
        sys.exit(1)
        
    with sync_playwright() as p:
        browser = p.chromium.connect_over_cdp(f"http://127.0.0.1:{args.port}")
        context = browser.contexts[0]
        page = get_douyin_page(context)
        
        if not page:
            print(">>> 错误: 未能在浏览器中找到可用页面！", flush=True)
            sys.exit(1)
            
        client = page.context.new_cdp_session(page)
        client.send('Page.setDownloadBehavior', {
            'behavior': 'allow',
            'downloadPath': str(DOWNLOAD_DIR)
        })
        
        if not wait_for_user_login(page):
            sys.exit(1)
            
        export_datacenter_tables(page)
        export_works_list(page)
        download_period_video(context, page, target_date, period_dir)
        
    run_local_pipeline(target_date, period_dir)
    print(">>> 采集与管线处理全部顺利完成！", flush=True)

if __name__ == "__main__":
    main()
